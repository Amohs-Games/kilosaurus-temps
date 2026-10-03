"""Tests de tools/migrate_d1.py : py -m unittest discover -s tests -p "test_*.py" """
import os
import sys
import tempfile
import unittest
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))
import migrate_d1 as m  # noqa: E402

HEADERS = ["ID", "Personne", "Date", "Début", "Fin", "Heures", "Note", "Source", "Saisi le", "Corrigé", "Modifié le", "Type"]


def workbook(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        if name == "_Config":
            for r, values in enumerate(rows, start=1):
                for c, v in enumerate(values, start=1):
                    ws.cell(r, c, v)
            continue
        ws.append(HEADERS)
        for row in rows:
            ws.append(row)
    wb.save(path)


def d(*args):
    return datetime(*args)


CONFIG = [["Personnes", "", "Paramètre", "Valeur"], ["Amohs", "", "Seuil oubli (h)", 8], ["", "", "Onglet studio", "Studio"], ["", "", "Boutons visibles", 4]]


class MigrateTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.path = os.path.join(self.dir, "export.xlsx")

    def run_build(self, sheets, exclude=()):
        workbook(self.path, sheets)
        cfg, projects, rows = m.read(self.path)
        return cfg, projects, m.build(cfg, projects, rows, set(exclude))

    def test_sessions_blocs_et_ordre_des_projets(self):
        cfg, projects, (entries, problems, warnings) = self.run_build({
            "Heirfall": [["s1", "Amohs", d(2026, 10, 1), d(2026, 10, 1, 9), d(2026, 10, 1, 12, 30), 3.5, "dialogues", "app", d(2026, 10, 1, 9), None, d(2026, 10, 1, 12, 30), "Narration"]],
            "Fluffy": [["report-0001", "Amohs", d(2024, 3, 2), None, None, 3.5, "menus", "report", None, None, None, None],
                       ["x1", "Alex", d(2026, 10, 1), None, None, 8, "", "app", d(2026, 10, 1, 9), None, None, "Dev"]],
            "_Config": CONFIG,
            "ZZ Test A1": [],
        })
        self.assertEqual(projects, ["Heirfall", "Fluffy"])
        self.assertEqual(cfg["me"], "Amohs")
        self.assertEqual(problems, [])
        self.assertEqual([e["id"] for e in entries], ["s1", "report-0001"])
        s1, rep = entries
        self.assertEqual(s1["kind"], "session")
        self.assertEqual(s1["type"], "Writing")
        self.assertEqual(s1["day"], "2026-10-01")
        self.assertEqual(s1["end_ms"] - s1["start_ms"], 3.5 * 3600000)
        self.assertEqual(rep["kind"], "block")
        self.assertEqual(rep["hours"], 3.5)
        self.assertEqual(rep["type"], "Misc")
        self.assertEqual(rep["day"], "2024-03-02")
        self.assertEqual(m.totals(entries), {"rows": 2, "byProjectType": {"Heirfall|Writing": 3.5, "Fluffy|Misc": 3.5}})

    def test_heure_de_paris_et_changement_d_heure(self):
        _, _, (entries, _, _) = self.run_build({
            "Fluffy": [["s1", "Amohs", d(2026, 10, 25), d(2026, 10, 25, 1), d(2026, 10, 25, 4), 4, "", "app", None, None, None, "Dev"]],
            "_Config": CONFIG,
        })
        # 01:00 (heure d'été) → 04:00 (heure d'hiver) le 25 octobre : 4 h réelles.
        self.assertEqual(entries[0]["end_ms"] - entries[0]["start_ms"], 4 * 3600000)
        self.assertEqual(entries[0]["start_ms"], int(datetime.fromisoformat("2026-10-24T23:00:00+00:00").timestamp() * 1000))

    def test_lignes_refusees(self):
        _, _, (entries, problems, _) = self.run_build({
            "Fluffy": [
                ["a", "Amohs", d(2026, 10, 1), d(2026, 10, 1, 9), d(2026, 10, 1, 8), 1, "", "app", None, None, None, "Dev"],
                ["b", "Amohs", d(2026, 10, 1), d(2026, 10, 1, 9), None, None, "", "app", None, None, None, "Dev"],
                ["c", "Amohs", d(2026, 9, 1), d(2026, 9, 1, 9), d(2026, 9, 2, 10), 25, "", "app", None, None, None, "Dev"],
                ["e", "Amohs", d(2026, 9, 5), None, None, 0, "", "app", None, None, None, "Dev"],
                ["e72", "Amohs", d(2026, 9, 5), None, None, 72, "", "report", None, None, None, None],
                ["f", "Amohs", d(2026, 9, 6), d(2026, 9, 6, 9), d(2026, 9, 6, 11), 2, "", "app", None, None, None, "Dev"],
                ["g", "Amohs", d(2026, 9, 6), d(2026, 9, 6, 10), d(2026, 9, 6, 12), 2, "", "app", None, None, None, "Dev"],
            ],
            "_Config": CONFIG,
        })
        ids = sorted(p["id"] for p in problems)
        self.assertEqual(ids, ["a", "b", "c", "e", "f", "g"])
        self.assertTrue(any("chevauche" in p["why"] for p in problems))

    def test_exclure_et_signaler_les_sessions_longues(self):
        sheets = {
            "Fluffy": [
                ["long", "Amohs", d(2026, 10, 2), d(2026, 10, 2, 13, 5), d(2026, 10, 3, 3, 7), 14.03, "", "app", None, None, None, "Dev"],
                ["ok", "Amohs", d(2026, 10, 1), d(2026, 10, 1, 9), d(2026, 10, 1, 10), 1, "", "app", None, None, None, "Dev"],
                ["stale", "Amohs", d(2026, 10, 1), d(2026, 10, 1, 10), d(2026, 10, 1, 11), 3, "", "app", None, None, None, "Dev"],
            ],
            "_Config": CONFIG,
        }
        _, _, (entries, problems, warnings) = self.run_build(sheets)
        self.assertEqual(problems, [])
        self.assertTrue(any(w["id"] == "long" and "8 h" in w["why"] for w in warnings))
        self.assertTrue(any(w["id"] == "stale" and "Heures" in w["why"] for w in warnings))
        _, _, (entries, _, warnings) = self.run_build(sheets, exclude=["long"])
        self.assertEqual([e["id"] for e in entries], ["ok", "stale"])
        self.assertFalse(any(w["id"] == "long" for w in warnings))

    def test_sql(self):
        cfg, projects, (entries, _, _) = self.run_build({
            "Fluffy": [["q'1", "Amohs", d(2026, 10, 1), None, None, 2, "l'écran", "app", d(2026, 10, 1, 9), None, None, "Dev"]],
            "_Config": CONFIG,
        })
        sql = m.to_sql(cfg, projects, entries)
        self.assertIn("DELETE FROM entries;", sql)
        self.assertIn("INSERT INTO projects (name, position) VALUES ('Fluffy', 1);", sql)
        self.assertIn("('threshold', '8')", sql)
        self.assertIn("'q''1'", sql)
        self.assertIn("'l''écran'", sql)


    def test_exclusions_et_autres_personnes_visibles(self):
        workbook(self.path, {
            "Fluffy": [
                ["a", "Amohs", d(2026, 10, 1), None, None, 2, "", "app", None, None, None, "Dev"],
                ["b", "Alex", d(2026, 10, 1), None, None, 2, "", "app", None, None, None, "Dev"],
                ["c", "Alex", d(2026, 10, 1), None, None, 2, "", "app", None, None, None, "Dev"],
            ],
            "_Config": CONFIG,
        })
        cfg, projects, rows = m.read(self.path)
        self.assertEqual(m.unknown_ids(rows, {"a", "zz"}), ["zz"])
        self.assertEqual(m.other_people(cfg, rows), {"Alex": 2})
        self.assertEqual(m.parse_args(["x.xlsx", "--exclude", "a,b", "--exclude", "c"]), ("x.xlsx", "reprise.sql", {"a", "b", "c"}))
        self.assertEqual(m.parse_args(["x.xlsx", "out.sql"]), ("x.xlsx", "out.sql", set()))

if __name__ == "__main__":
    unittest.main()
