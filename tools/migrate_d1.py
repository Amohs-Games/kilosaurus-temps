"""Reprise de l'historique : la Google Sheet (téléchargée en .xlsx) → un fichier SQL pour D1.

Usage : py tools/migrate_d1.py export.xlsx [reprise.sql] [--exclude id1,id2]

Lit les onglets de projet (tout onglet dont le nom ne commence ni par « _ » ni par « ZZ Test », dans
l'ordre des onglets), _Config, et les lignes de la personne de _Config. Refuse de produire le
fichier si une ligne viole une règle de la base (fin avant début, session de plus de 24 h,
chevauchement, bloc sans durée positive, compteur en cours) : chacune se corrige dans la Sheet, ou s'écarte avec --exclude.
Signale sans bloquer les sessions plus longues que le seuil d'oubli et les lignes dont la colonne
Heures ne correspond plus à Début / Fin. Écrit aussi <sortie>.totaux.json, que relit
tools/check_import.mjs.
"""
import json
import sys
from collections import defaultdict
from datetime import date, datetime, time
from zoneinfo import ZoneInfo

import openpyxl

PARIS = ZoneInfo("Europe/Paris")
LEGACY = {"Autre": "Misc", "Narration": "Writing"}
TYPES = ["Misc", "Writing", "Sound", "Market", "Concept", "Art", "TechArt", "UI", "Dev", "Gameplay", "Tooling", "Debug", "Pause"]
COLS = ["id", "person", "date", "start", "end", "hours", "note", "source", "created", "corrected", "touched", "type"]
ENTRY = ["id", "project", "kind", "day", "start_ms", "end_ms", "hours", "type", "note", "source", "created_ms", "touched_ms", "deleted_ms"]
HOUR = 3600000
MAX_MS = 24 * HOUR


def to_ms(value):
    """Date-heure de la Sheet (heure de Paris, sans fuseau) → millisecondes UTC."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        dt = value if value.tzinfo else value.replace(tzinfo=PARIS)
        return int(dt.timestamp() * 1000)
    if isinstance(value, date):
        return to_ms(datetime.combine(value, time()))
    raise ValueError(f"date illisible : {value!r}")


def day_key(ms):
    return datetime.fromtimestamp(ms / 1000, PARIS).strftime("%Y-%m-%d")


def read_config(ws):
    persons = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    params = {ws.cell(r, 3).value: ws.cell(r, 4).value for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value}
    return {
        "me": persons[0],
        "threshold": params.get("Seuil oubli (h)", 8),
        "studio": params.get("Onglet studio", "Studio"),
        "visible": params.get("Boutons visibles", 4),
    }


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    cfg = read_config(wb["_Config"])
    projects = [ws.title for ws in wb.worksheets if not ws.title.startswith("_") and not ws.title.startswith("ZZ Test")]
    rows = []
    for name in projects:
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            v = dict(zip(COLS, (ws.cell(r, c).value for c in range(1, len(COLS) + 1))))
            if not v["id"] or not v["person"]:
                continue
            v["project"] = name
            v["where"] = f"{name}, ligne {r}"
            rows.append(v)
    return cfg, projects, rows


def convert(v):
    """Une ligne de la Sheet → (ligne de la base, None) ou (None, raison du refus)."""
    start, end = to_ms(v["start"]), to_ms(v["end"])
    created = to_ms(v["created"]) or start or to_ms(v["date"])
    common = {
        "id": str(v["id"]), "project": v["project"], "type": LEGACY.get(v["type"] or "Misc", v["type"] or "Misc"),
        "note": str(v["note"] or "").strip()[:500], "source": v["source"] or "app",
        "created_ms": created, "touched_ms": to_ms(v["touched"]) or created, "deleted_ms": None,
    }
    if common["type"] not in TYPES:
        common["type"] = "Misc"
    if start is not None:
        if end is None:
            return None, "compteur en cours : arrête-le avant la reprise"
        if end < start:
            return None, "fin avant le début"
        if end - start > MAX_MS:
            return None, "session de plus de 24 h"
        return {**common, "kind": "session", "day": day_key(start), "start_ms": start, "end_ms": end, "hours": None}, None
    hours = v["hours"]
    if not isinstance(hours, (int, float)) or hours <= 0:
        return None, f"bloc avec une durée invalide ({hours!r})"
    if v["date"] is None:
        return None, "bloc sans date"
    return {**common, "kind": "block", "day": day_key(to_ms(v["date"])), "start_ms": None, "end_ms": None, "hours": float(hours)}, None


def build(cfg, projects, rows, exclude):
    """→ (lignes de la base, refus, signalements). Un refus empêche la reprise ; un signalement non."""
    entries, problems, warnings = [], [], []
    threshold_ms = float(cfg["threshold"]) * HOUR
    for v in rows:
        if v["person"] != cfg["me"] or str(v["id"]) in exclude:
            continue
        e, why = convert(v)
        if why:
            problems.append({"id": str(v["id"]), "where": v["where"], "why": why})
            continue
        if e["kind"] == "session":
            length = e["end_ms"] - e["start_ms"]
            if length > threshold_ms:
                warnings.append({"id": e["id"], "where": v["where"], "why": f"session de {length / HOUR:.2f} h (plus de {cfg['threshold']} h)"})
            if isinstance(v["hours"], (int, float)) and abs(v["hours"] - length / HOUR) > 0.01:
                warnings.append({"id": e["id"], "where": v["where"], "why": f"Heures = {v['hours']} mais Début / Fin = {length / HOUR:.2f} h (la base garde Début / Fin)"})
        entries.append(e)
    sessions = sorted((e for e in entries if e["kind"] == "session"), key=lambda e: e["start_ms"])
    for a, b in zip(sessions, sessions[1:]):
        if b["start_ms"] < a["end_ms"]:
            why = f"chevauche {a['id']} ({a['project']}, {day_key(a['start_ms'])})"
            problems.append({"id": b["id"], "where": b["project"], "why": why})
            problems.append({"id": a["id"], "where": a["project"], "why": f"chevauche {b['id']}"})
    return entries, problems, warnings


def lit(value):
    if value is None:
        return "NULL"
    if isinstance(value, (int, float)):
        return repr(value)
    return "'" + str(value).replace("'", "''") + "'"


def to_sql(cfg, projects, entries):
    out = ["-- Reprise de l'historique (tools/migrate_d1.py). Remplace tout le contenu de la base.",
           "DELETE FROM entries;", "DELETE FROM projects;", "DELETE FROM settings;", "DELETE FROM ops;"]
    for i, name in enumerate(projects, start=1):
        out.append(f"INSERT INTO projects (name, position) VALUES ({lit(name)}, {i});")
    for key in ("me", "studio", "visible", "threshold"):
        out.append(f"INSERT INTO settings (key, value) VALUES ({lit(key)}, {lit(str(cfg[key]))});")
    for e in entries:
        out.append(f"INSERT INTO entries ({', '.join(ENTRY)}) VALUES ({', '.join(lit(e[c]) for c in ENTRY)});")
    return "\n".join(out) + "\n"


def totals(entries):
    by = defaultdict(float)
    for e in entries:
        by[f"{e['project']}|{e['type']}"] += e["hours"] if e["kind"] == "block" else (e["end_ms"] - e["start_ms"]) / HOUR
    return {"rows": len(entries), "byProjectType": {k: round(v, 4) for k, v in by.items()}}


def parse_args(argv):
    """→ (source, sortie, ids à écarter). Plusieurs --exclude s'ajoutent."""
    exclude, args, i = set(), [], 0
    while i < len(argv):
        if argv[i] == "--exclude":
            exclude |= {x for x in argv[i + 1].split(",") if x}
            i += 2
            continue
        args.append(argv[i])
        i += 1
    return args[0], (args[1] if len(args) > 1 else "reprise.sql"), exclude


def unknown_ids(rows, exclude):
    """Ids demandés à l'écart mais absents de la Sheet : une faute de frappe ferait reprendre la ligne."""
    known = {str(v["id"]) for v in rows}
    return sorted(x for x in exclude if x not in known)


def other_people(cfg, rows):
    """Lignes des autres personnes de la Sheet, non reprises : comptées pour qu'aucune perte ne passe inaperçue."""
    count = defaultdict(int)
    for v in rows:
        if v["person"] != cfg["me"]:
            count[v["person"]] += 1
    return dict(count)


def main(argv):
    sys.stdout.reconfigure(encoding="utf-8")  # accents lisibles dans la console Windows
    src, out, exclude = parse_args(argv)
    cfg, projects, rows = read(src)
    unknown = unknown_ids(rows, exclude)
    if unknown:
        print(f"ÉCHEC : ids à écarter introuvables dans la Sheet : {', '.join(unknown)}")
        raise SystemExit(1)
    entries, problems, warnings = build(cfg, projects, rows, exclude)

    print(f"Projets : {', '.join(projects)}")
    print(f"{len(entries)} lignes reprises pour {cfg['me']} ({len(exclude)} écartée(s) à la demande : {', '.join(sorted(exclude)) or 'aucune'}).")
    for person, n in sorted(other_people(cfg, rows).items()):
        print(f"  non reprises : {n} ligne(s) de {person}")
    for w in warnings:
        print(f"  à vérifier : {w['id']} ({w['where']}) : {w['why']}")
    if problems:
        print("ÉCHEC : lignes refusées par les règles de la base (corrige-les dans la Sheet, ou --exclude <id>) :")
        for p in problems:
            print(f"  {p['id']} ({p['where']}) : {p['why']}")
        raise SystemExit(1)
    t = totals(entries)
    by_project = defaultdict(float)
    for key, hours in t["byProjectType"].items():
        by_project[key.split("|")[0]] += hours
    for name in projects:
        print(f"  {name:12} {by_project.get(name, 0):9.2f} h")
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write(to_sql(cfg, projects, entries))
    with open(out.rsplit(".", 1)[0] + ".totaux.json", "w", encoding="utf-8") as f:
        json.dump(t, f, ensure_ascii=False, indent=1)
    print(f"OK : {out} écrit. Vérifie-le avec : node tools/check_import.mjs {out}")


if __name__ == "__main__":
    main(sys.argv[1:])
